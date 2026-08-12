"""Reference-data export — the inverse of app.importer.import_service.

Loads the current server-side template workbook as a structural base
(preserves banners/headers/formatting exactly, discovered dynamically from
each sheet's own header row rather than hardcoded) and overwrites each
sheet's data rows with live DB content, so "download, edit, re-upload"
round-trips through the same importer without this code needing to
reconstruct sheet layout from scratch.

Every sheet here is upsert-only on import EXCEPT VISA RULES, which does a
full delete-then-reload with no natural key (see loaders/visa.py) — the
export MUST include every current VisaRule row with country/nationality
names that resolve identically, or a re-upload silently deletes whatever's
missing. See _export_visa_rules.

Known, deliberate lossy points (the DB holds more than this sheet format
can express — exporting always writes the best available inverse, never
guesses at the rest):
  - Client: 4 address columns (billing_address_1/2, city, country) collapse
    to one DB column; exported entirely into billing_address_1, the other
    three left blank.
  - Vendor.messaging_cc and .questionnaire_answers, VendorCoverageAirport
    .contacts: no sheet column exists for these at all.
  - User.operator_scope/client_scope: NULL and explicit "ALL" are not
    distinguished in the DB; NULL always exports as "ALL".
  - VisaRule.source/verified_by/verified_on/expires_on: no sheet column.
  - Airport: "Operating Hours" and "Airport Hours" both feed one DB column
    on import; exported only into "Operating Hours".
"""

import io

import openpyxl
from openpyxl.workbook.workbook import Workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import get_settings
from app.models.aircraft import Aircraft, AircraftPerformance
from app.models.airport import Airport
from app.models.client import Client
from app.models.country import Country, GroundHandlingPolicy
from app.models.country_requirements import CountryRequirement
from app.models.messaging import MessageTemplate
from app.models.operator import Operator
from app.models.service_catalogue import ServiceCatalogueEntry, ServiceLevel
from app.models.settings import Setting
from app.models.user import User
from app.models.vendor import Vendor, VendorCoverageAirport, VendorCoverageCountry
from app.models.visa import VisaMatrixCell, VisaRule


def _column_index_by_header(ws, header_row: int) -> dict[str, int]:
    idx: dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=header_row, column=c).value
        if v is not None:
            idx[str(v).strip()] = c
    return idx


def _write_simple_sheet(ws, *, header_row: int, data_start_row: int, rows: list[dict]) -> None:
    """For sheets with nothing else below the data table — clears from
    data_start_row through the greater of the sheet's current extent or
    the new row count, then writes. openpyxl auto-grows the sheet if
    `rows` needs more rows than currently exist.
    """
    col_idx = _column_index_by_header(ws, header_row)
    clear_through = max(ws.max_row, data_start_row + len(rows) - 1) if rows else ws.max_row
    for r in range(data_start_row, clear_through + 1):
        for c in range(1, ws.max_column + 1):
            ws.cell(row=r, column=c).value = None
    for i, row in enumerate(rows):
        r = data_start_row + i
        for header, value in row.items():
            c = col_idx.get(header)
            if c is not None:
                ws.cell(row=r, column=c).value = value


def _write_capped_sheet(ws, *, header_row: int, data_start_row: int, rows: list[dict]) -> None:
    """For sheets with an unrelated section directly below the data table
    (MESSAGE TEMPLATES, USERS & SETTINGS) — finds the table's current
    capacity (first fully-blank row), inserts extra blank rows before that
    boundary if `rows` needs more room, then clears and writes. Paired
    with loaders/messaging.py and loaders/users.py's stop_at_blank=True
    (see app/importer/xlsx_reader.py) so a re-upload reads back however
    many rows were written here, not a hardcoded count.
    """
    col_idx = _column_index_by_header(ws, header_row)
    r = data_start_row
    while True:
        values = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        if all(v is None for v in values):
            break
        r += 1
    capacity = r - data_start_row
    if len(rows) > capacity:
        ws.insert_rows(r, amount=len(rows) - capacity)
    window = max(len(rows), capacity)
    for rr in range(data_start_row, data_start_row + window):
        for c in range(1, ws.max_column + 1):
            ws.cell(row=rr, column=c).value = None
    for i, row in enumerate(rows):
        rr = data_start_row + i
        for header, value in row.items():
            c = col_idx.get(header)
            if c is not None:
                ws.cell(row=rr, column=c).value = value


def _yn(value: bool | None) -> str | None:
    if value is None:
        return None
    return "YES" if value else "NO"


def _join(values: list[str] | None) -> str | None:
    return ", ".join(values) if values else None


async def _export_countries(session: AsyncSession, ws) -> dict[str, Country]:
    countries = (await session.execute(select(Country).where(Country.deleted_at.is_(None)))).scalars().all()
    rows = []
    for c in countries:
        mandatory = self_handling = None
        if c.ground_handling_policy == GroundHandlingPolicy.MANDATORY:
            mandatory, self_handling = "YES", "NO"
        elif c.ground_handling_policy == GroundHandlingPolicy.SELF_HANDLING_PERMITTED:
            mandatory, self_handling = "NO", "YES"
        elif c.ground_handling_policy == GroundHandlingPolicy.NOT_MANDATORY:
            mandatory, self_handling = "NO", "NO"
        rows.append(
            {
                "iso_alpha3": c.iso3,
                "country_name": c.name,
                "iso_alpha2": c.iso2,
                "overflight_permit_required": "REQUIRED" if c.overflight_permit_required else "NOT REQUIRED",
                "landing_permit_required": "REQUIRED" if c.landing_permit_required else "NOT REQUIRED",
                "ground_handling_mandatory": mandatory,
                "self_handling_permitted": self_handling,
                "calling_code": c.calling_code,
                "civil_aviation_authority": c.caa_name,
                "region": c.region,
                "permit_flags_source": c.permit_flags_source,
                "permit_flags_verified_by": c.permit_flags_verified_by,
                "permit_flags_verified_on": c.permit_flags_verified_on,
                "standard_lead_time_hours": c.standard_lead_time_hours,
                "lead_time_source": c.lead_time_source,
                "lead_time_verified_by": c.lead_time_verified_by,
                "lead_time_verified_on": c.lead_time_verified_on,
            }
        )
    rows.sort(key=lambda r: r["iso_alpha3"])
    _write_simple_sheet(ws, header_row=1, data_start_row=2, rows=rows)
    return {c.iso3: c for c in countries}


async def _export_airports(session: AsyncSession, ws, countries_by_iso3: dict[str, Country]) -> None:
    airports = (await session.execute(select(Airport).where(Airport.deleted_at.is_(None)))).scalars().all()
    rows = [
        {
            "Airport ICAO": a.icao,
            "Latitude (DD)": a.lat,
            "Longitude (DD)": a.lon,
            "Country": countries_by_iso3[a.country_iso3].name if a.country_iso3 in countries_by_iso3 else None,
            "Operating Hours": a.operating_hours,
            "Airport IATA": a.iata,
            "Airport Name": a.name,
            "City Name": a.city,
            "Elevation (ft)": a.elevation_ft,
            "Time Zone": a.timezone,
            "Airport Type": a.airport_type,
            "Airport of Entry": _yn(a.is_airport_of_entry),
            "PPR Required": _yn(a.ppr_required),
            "Fuel Types Available": _join(a.fuel_grades),
        }
        for a in sorted(airports, key=lambda a: a.icao)
    ]
    _write_simple_sheet(ws, header_row=2, data_start_row=3, rows=rows)


async def _export_operators(session: AsyncSession, ws) -> dict:
    operators = (await session.execute(select(Operator).where(Operator.deleted_at.is_(None)))).scalars().all()
    rows = [
        {
            "operator_id": o.source_ref,
            "operator_name": o.name,
            "status": o.status.value,
            "aoc_number": o.aoc_number,
            "icao_designator": o.icao_designator,
            "iata_designator": o.iata_designator,
            "home_base_icao": o.home_base_icao,
            "primary_contact_name": o.contact_name,
            "primary_contact_phone": o.contact_phone,
            "occ_email": o.occ_email,
            "billing_email": o.billing_email,
            "default_currency": o.currency,
            "tax_id": o.tax_id,
            "preferred_channel": o.preferred_channel.value if o.preferred_channel else None,
            "messaging_to_emails": _join(o.messaging_to),
            "messaging_cc_emails": _join(o.messaging_cc),
            "sita_address": o.sita_address,
            "aftn_address": o.aftn_address,
            "message_format": o.message_format.value if o.message_format else None,
            "sending_team_signature": o.sending_team_signature,
        }
        for o in operators
        if o.source_ref
    ]
    rows.sort(key=lambda r: r["operator_id"])
    _write_simple_sheet(ws, header_row=1, data_start_row=2, rows=rows)
    return {o.id: o.source_ref for o in operators if o.source_ref}


async def _export_aircraft_performance(session: AsyncSession, ws) -> None:
    perf = (await session.execute(select(AircraftPerformance))).scalars().all()
    rows = [
        {
            "icao_type": p.icao_type,
            "manufacturer": p.manufacturer,
            "model_series": p.model_series,
            "max_range_nm": p.max_range_nm,
            "cruise_tas_kts": p.cruise_tas_kts,
            "fuel_burn_kg_hr": p.fuel_burn_kg_per_hr,
            "max_pax": p.max_pax,
            "mtow_kg": p.mtow_kg,
            "source": p.source,
            "verified": _yn(p.verified),
            "verified_by": p.verified_by,
            "verified_on": p.verified_on,
        }
        for p in sorted(perf, key=lambda p: p.icao_type)
    ]
    _write_simple_sheet(ws, header_row=6, data_start_row=7, rows=rows)


async def _export_aircraft(session: AsyncSession, ws, operator_source_ref_by_id: dict) -> None:
    aircraft = (await session.execute(select(Aircraft))).scalars().all()
    rows = [
        {
            "registration": a.registration,
            "icao_type": a.icao_type,
            "operator_ref": operator_source_ref_by_id.get(a.operator_id),
            "aircraft_status": a.status.value,
            "manufacturer": a.manufacturer,
            "model_series": a.model_series,
            "serial_number": a.serial_number,
            "home_base_icao": a.home_base_icao,
            "mtow_kg": a.mtow_kg,
            "max_pax": a.max_pax,
            "cof_expiry_date": a.cofa_expiry,
            "insurance_expiry_date": a.insurance_expiry,
            "total_airframe_hours": a.total_hours,
            "total_landings": a.total_landings,
        }
        for a in sorted(aircraft, key=lambda a: a.registration)
    ]
    _write_simple_sheet(ws, header_row=1, data_start_row=2, rows=rows)


async def _export_clients(session: AsyncSession, ws, operator_source_ref_by_id: dict) -> dict:
    clients = (await session.execute(select(Client).where(Client.deleted_at.is_(None)))).scalars().all()
    rows = []
    for c in clients:
        if not c.source_ref:
            continue
        credit_limit = c.credit_limit_minor_units / 100 if c.credit_limit_minor_units is not None else None
        rows.append(
            {
                "client_id": c.source_ref,
                "operator_id": operator_source_ref_by_id.get(c.operator_id),
                "bill_to_legal_name": c.bill_to_legal_name,
                "client_name": c.bill_to_legal_name,
                "billing_address_1": c.address,
                "tax_vat_id": c.tax_vat_number,
                "billing_contact_name": c.billing_contact_name,
                "billing_email": c.billing_contact_email,
                "currency": c.currency,
                "payment_terms": c.payment_terms,
                "credit_limit": credit_limit,
                "preferred_channel": c.preferred_channel.value if c.preferred_channel else None,
                "messaging_to_emails": _join(c.messaging_to),
                "messaging_cc_emails": _join(c.messaging_cc),
                "sita_address": c.sita_address,
                "aftn_address": c.aftn_address,
                "message_format": c.message_format.value if c.message_format else None,
                "sending_team_signature": c.sending_team_signature,
            }
        )
    rows.sort(key=lambda r: r["client_id"])
    _write_simple_sheet(ws, header_row=2, data_start_row=3, rows=rows)
    return {c.id: c.source_ref for c in clients if c.source_ref}


async def _export_vendors(session: AsyncSession, ws) -> dict:
    vendors = (await session.execute(select(Vendor).where(Vendor.deleted_at.is_(None)))).scalars().all()
    rows = [
        {
            "vendor_id": v.source_ref,
            "company_name": v.name,
            "email": v.messaging_to[0] if v.messaging_to else None,
            "services_offered_codes": _join(v.service_scope),
            "sita_address": v.sita_address,
            "aftn_address": v.aftn_address,
            "questionnaire_sent_on": v.questionnaire_sent_on,
            "questionnaire_returned_on": v.questionnaire_returned_on,
            "capability_status": v.capability_status.value,
            "approved_by": v.approved_by,
            "approved_on": v.approved_on,
            "preference_rank": v.preference_rank,
        }
        for v in vendors
        if v.source_ref
    ]
    rows.sort(key=lambda r: r["vendor_id"])
    _write_simple_sheet(ws, header_row=1, data_start_row=2, rows=rows)
    return {v.id: v.source_ref for v in vendors if v.source_ref}


async def _export_vendor_coverage_airports(session: AsyncSession, ws, vendor_source_ref_by_id: dict) -> None:
    rows_db = (await session.execute(select(VendorCoverageAirport))).scalars().all()
    rows = [
        {
            "id": r.source_ref,
            "vendor_id": vendor_source_ref_by_id.get(r.vendor_id),
            "icao_code": r.icao,
            "is_primary_handler": _yn(r.is_primary_handler),
            "fbo_facility_name": r.fbo_name,
            "sita_address": r.sita,
            "aftn_address": r.aftn,
            "vhf_frequency": r.vhf,
        }
        for r in rows_db
        if r.source_ref
    ]
    rows.sort(key=lambda r: r["id"])
    _write_simple_sheet(ws, header_row=1, data_start_row=2, rows=rows)


async def _export_vendor_coverage_countries(
    session: AsyncSession, ws, vendor_source_ref_by_id: dict, countries_by_iso3: dict[str, Country]
) -> None:
    rows_db = (await session.execute(select(VendorCoverageCountry))).scalars().all()
    rows = [
        {
            "id": r.source_ref,
            "vendor_id": vendor_source_ref_by_id.get(r.vendor_id),
            "country_iso_code": countries_by_iso3[r.country_iso3].iso2 if r.country_iso3 in countries_by_iso3 else None,
            "has_caa_direct_account": _yn(r.has_caa_direct_account),
            "service_scope": _join(r.service_scope),
        }
        for r in rows_db
        if r.source_ref
    ]
    rows.sort(key=lambda r: r["id"])
    _write_simple_sheet(ws, header_row=1, data_start_row=2, rows=rows)


async def _export_service_catalogue(session: AsyncSession, ws) -> None:
    entries = (await session.execute(select(ServiceCatalogueEntry))).scalars().all()
    code_by_id = {e.id: e.code for e in entries}
    services = [e for e in entries if e.level == ServiceLevel.SERVICE]
    sub_services = [e for e in entries if e.level == ServiceLevel.SUB_SERVICE]

    rows = []
    for e in sorted(services, key=lambda e: e.code):
        rows.append(
            {
                "level": "SERVICE",
                "id": e.code,
                "code": e.code,
                "category": e.category.value,
                "name": e.name,
                "description": e.description,
                "ground_grid_order": e.ground_grid_order,
            }
        )
    for e in sorted(sub_services, key=lambda e: e.code):
        parent_code = code_by_id.get(e.parent_service_id) if e.parent_service_id else None
        rows.append(
            {
                "level": "SUB-SERVICE",
                "id": e.code,
                "code": e.code,
                "name": e.name,
                "parent_service_id": parent_code,
                "description": e.description,
            }
        )
    _write_simple_sheet(ws, header_row=2, data_start_row=3, rows=rows)


async def _export_visa_matrix(session: AsyncSession, ws, countries_by_iso3: dict[str, Country]) -> None:
    """Only updates cells already represented by the template's existing
    row/column pairs — this sheet ships "deliberately near-empty" by
    design (see loaders/visa.py); export does not expand the grid to cover
    countries the template doesn't already list as a destination row or
    nationality column.
    """
    cells = {
        (c.country_iso3, c.nationality_iso3): c
        for c in (await session.execute(select(VisaMatrixCell))).scalars().all()
    }
    header_row = 3
    nationality_cols: list[tuple[int, str]] = []
    for c in range(2, ws.max_column + 1):
        name = ws.cell(row=header_row, column=c).value
        if not name:
            continue
        iso3 = next((iso3 for iso3, country in countries_by_iso3.items() if country.name.strip().upper() == str(name).strip().upper()), None)
        if iso3:
            nationality_cols.append((c, iso3))

    for r in range(4, ws.max_row + 1):
        dest_name = ws.cell(row=r, column=1).value
        if not dest_name:
            continue
        dest_iso3 = next(
            (iso3 for iso3, country in countries_by_iso3.items() if country.name.strip().upper() == str(dest_name).strip().upper()),
            None,
        )
        if not dest_iso3:
            continue
        for col, nat_iso3 in nationality_cols:
            cell = cells.get((dest_iso3, nat_iso3))
            ws.cell(row=r, column=col).value = cell.requirement.value if cell else None


async def _export_visa_rules(session: AsyncSession, ws, countries_by_iso3: dict[str, Country]) -> None:
    """CRITICAL: VISA RULES does a full delete-then-reload on import with
    no natural key (see loaders/visa.py) — every current row must appear
    here with names that resolve identically, or a re-upload deletes it.
    """
    rules = (await session.execute(select(VisaRule))).scalars().all()
    rows = [
        {
            "Required?": r.requirement.value,
            "Country": countries_by_iso3[r.country_iso3].name if r.country_iso3 in countries_by_iso3 else None,
            "Nationality": countries_by_iso3[r.nationality_iso3].name if r.nationality_iso3 in countries_by_iso3 else None,
            "Airport ICAO (optional)": r.airport_icao,
            "Visa Type": r.visa_type,
            "Lead Time (days)": r.lead_time_days,
            "Max Stay (days)": r.max_stay_days,
            "Notes / Conditions": r.conditions,
        }
        for r in rules
    ]
    _write_simple_sheet(ws, header_row=1, data_start_row=2, rows=rows)


async def _export_country_requirements(session: AsyncSession, ws, countries_by_iso3: dict[str, Country]) -> None:
    reqs = (
        await session.execute(select(CountryRequirement).where(CountryRequirement.deleted_at.is_(None)))
    ).scalars().all()
    rows = [
        {
            "req_id": r.source_ref,
            "request_type": r.request_type,
            "country": countries_by_iso3[r.country_iso3].name if r.country_iso3 in countries_by_iso3 else None,
            "forms_required": _join(r.required_forms),
            "documents_required": _join(r.required_documents),
            "lead_time_hours": r.lead_time_hours,
            "authority_working_hours": r.authority_working_hours,
            "accepted_channels": _join(r.accepted_channels),
            "local_agent_required": _yn(r.local_agent_required),
            "caa_direct_account_required": _yn(r.direct_caa_account_possible),
            "known_quirks_intel": r.known_quirks,
            "source": r.source,
            "verified_by": r.verified_by,
            "verified_on": r.verified_on,
            "review_due": r.review_due,
        }
        for r in reqs
        if r.source_ref
    ]
    rows.sort(key=lambda r: r["req_id"])
    _write_simple_sheet(ws, header_row=4, data_start_row=5, rows=rows)


async def _export_message_templates(session: AsyncSession, ws) -> None:
    templates = (await session.execute(select(MessageTemplate))).scalars().all()
    rows = [
        {
            "template_id": t.template_key,
            "template_name": t.name,
            "message_type": t.message_type,
            "recipient_role": t.recipient_role,
            "channel": t.channel.value,
            "body_template": t.body,
            "subject_line": t.subject_line,
            "footer_block": t.footer_block,
            "active": _yn(t.active),
        }
        for t in sorted(templates, key=lambda t: t.template_key)
    ]
    _write_capped_sheet(ws, header_row=37, data_start_row=38, rows=rows)


async def _export_users(session: AsyncSession, ws, operator_source_ref_by_id: dict, client_source_ref_by_id: dict) -> None:
    users = (await session.execute(select(User))).scalars().all()
    rows = [
        {
            "email": u.email,
            "full_name": u.full_name,
            "role": u.role.value.replace("_", " "),
            "operator_scope": operator_source_ref_by_id.get(u.operator_id, "ALL") if u.operator_id else "ALL",
            "client_scope": client_source_ref_by_id.get(u.client_id, "ALL") if u.client_id else "ALL",
            "status": "ACTIVE" if u.is_active else "INACTIVE",
            "mfa_enabled": _yn(u.mfa_enabled),
        }
        for u in sorted(users, key=lambda u: u.email)
    ]
    _write_capped_sheet(ws, header_row=14, data_start_row=15, rows=rows)


_SETTINGS_ROW_TO_KEY = {
    15: "default_permit_lead_time_hours",
    16: "default_ground_notice_hours",
    17: "range_reserve_margin",
    18: "passport_validity_buffer_days",
    19: "document_expiry_alert_days",
    20: "visa_rule_staleness_days",
    21: "allow_unverified_for_planning",
}


async def _export_admin_settings(session: AsyncSession, ws) -> None:
    settings_by_key = {s.key: s for s in (await session.execute(select(Setting))).scalars().all()}
    for row_num, key in _SETTINGS_ROW_TO_KEY.items():
        setting = settings_by_key.get(key)
        if setting is not None:
            ws.cell(row=row_num, column=2).value = setting.value


async def export_reference_workbook(session: AsyncSession) -> bytes:
    settings = get_settings()
    wb: Workbook = openpyxl.load_workbook(settings.import_source_xlsx)

    countries_by_iso3 = await _export_countries(session, wb["countries"])
    await _export_airports(session, wb["airports"], countries_by_iso3)
    operator_source_ref_by_id = await _export_operators(session, wb["Operators"])
    await _export_aircraft_performance(session, wb["AIRCRAFT PERFORMANCE"])
    await _export_aircraft(session, wb["Aircraft"], operator_source_ref_by_id)
    client_source_ref_by_id = await _export_clients(session, wb["Clients"], operator_source_ref_by_id)
    vendor_source_ref_by_id = await _export_vendors(session, wb["Vendors"])
    await _export_vendor_coverage_airports(session, wb["vendor_coverage_airports"], vendor_source_ref_by_id)
    await _export_vendor_coverage_countries(session, wb["vendor_coverage_countries"], vendor_source_ref_by_id, countries_by_iso3)
    await _export_service_catalogue(session, wb["SERVICE CATALOGUE"])
    await _export_visa_matrix(session, wb["VISA MATRIX"], countries_by_iso3)
    await _export_visa_rules(session, wb["VISA RULES"], countries_by_iso3)
    await _export_message_templates(session, wb["MESSAGE TEMPLATES"])
    await _export_country_requirements(session, wb["COUNTRY REQUIREMENTS"], countries_by_iso3)
    await _export_users(session, wb["USERS & SETTINGS"], operator_source_ref_by_id, client_source_ref_by_id)
    await _export_admin_settings(session, wb["admin"])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
