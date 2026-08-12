"""Thin helpers over openpyxl. The workbook has section-header rows and
prose banners above most data tables, so each loader states exactly which
row the column headers live on and which row data starts — there is no
attempt to auto-detect a header row, that is exactly the kind of "clever"
guess the verification contract forbids.
"""

from collections.abc import Iterator
from datetime import date, datetime
from typing import Any

import openpyxl
from openpyxl.workbook.workbook import Workbook


def load_workbook(path: str) -> Workbook:
    return openpyxl.load_workbook(path, data_only=True)


def read_rows(
    ws,
    *,
    header_row: int,
    data_start_row: int,
    data_end_row: int | None = None,
    max_col: int | None = None,
    stop_at_blank: bool = False,
) -> Iterator[dict[str, Any]]:
    """Yield one dict per data row keyed by the exact header cell text.

    Rows that are entirely blank are normally skipped (not yielded) but the
    scan continues to `data_end_row`/the sheet's last row. `stop_at_blank`
    instead ends the scan at the first entirely-blank row — for sheets with
    no `data_end_row` boundary where an unrelated section (a banner, a
    different table) sits directly below the data table separated by a
    blank row, e.g. MESSAGE TEMPLATES/USERS & SETTINGS: without this, a
    data_end_row bound had to be hand-picked and capped the sheet at a
    fixed row count; with it, the sheet can hold as many rows as an export
    inserts and still stop exactly at the real end of the table.
    """
    last_col = max_col or ws.max_column
    headers = [ws.cell(row=header_row, column=c).value for c in range(1, last_col + 1)]
    end = data_end_row or ws.max_row
    for r in range(data_start_row, end + 1):
        values = [ws.cell(row=r, column=c).value for c in range(1, last_col + 1)]
        if all(v is None for v in values):
            if stop_at_blank:
                break
            continue
        row = {}
        for header, value in zip(headers, values):
            if header is None:
                continue
            row[str(header).strip()] = value
        yield row


def s(row: dict, key: str) -> str | None:
    """String cell, trimmed, empty string treated as None."""
    v = row.get(key)
    if v is None:
        return None
    v = str(v).strip()
    return v or None


def f(row: dict, key: str) -> float | None:
    v = row.get(key)
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def i(row: dict, key: str) -> int | None:
    v = f(row, key)
    return int(v) if v is not None else None


def b(row: dict, key: str, *, true_values=("YES", "TRUE", "1")) -> bool | None:
    v = s(row, key)
    if v is None:
        return None
    return v.strip().upper() in true_values


def d(row: dict, key: str) -> date | None:
    v = row.get(key)
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def code(row: dict, key: str, *, max_len: int) -> str | None:
    """A short fixed-format code (ICAO, IATA, calling code, ...). Returns
    None rather than raising when the source cell holds something else
    (a stray country name, a note) — messy input data, not a schema
    violation to fail loudly on.
    """
    v = s(row, key)
    if v is None or len(v) > max_len:
        return None
    return v


def list_(row: dict, key: str, *, sep: str = ",") -> list[str] | None:
    v = s(row, key)
    if v is None:
        return None
    items = [part.strip() for part in v.replace(";", sep).split(sep) if part.strip()]
    return items or None
